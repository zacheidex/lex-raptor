"""Bounded, local extraction. No OCR, formula evaluation, macros, or network."""

import re
import zipfile
from dataclasses import asdict, dataclass, field
from email import policy
from email.parser import BytesParser
from hashlib import sha256
from io import BytesIO
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from .settings import settings

VERSION = "extract-1.0"
SUPPORTED = {".pdf", ".docx", ".txt", ".eml", ".xlsx"}


class IngestionError(ValueError):
    pass


@dataclass
class Chunk:
    id: str
    document_hash: str
    document_name: str
    locator: str
    text: str


@dataclass
class Extraction:
    name: str
    sha256: str
    version: str = VERSION
    status: str = "ready"
    chunks: list[Chunk] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    error_code: str | None = None

    def to_dict(self):
        return asdict(self)


def check_package(raw):
    with zipfile.ZipFile(BytesIO(raw)) as z:
        infos = z.infolist()
        if (
            len(infos) > 10000
            or sum(i.file_size for i in infos) > settings.max_expanded_bytes
        ):
            raise IngestionError("decompressed_size_limit")
        for i in infos:
            if (
                i.flag_bits & 1
                or ".." in Path(i.filename).parts
                or i.filename.startswith("/")
            ):
                raise IngestionError("unsafe_archive")
            if i.file_size > settings.max_expanded_bytes:
                raise IngestionError("decompressed_size_limit")
        if any("vbaproject" in n.lower() for n in z.namelist()):
            raise IngestionError("macros_not_supported")


def extract(name: str, raw: bytes) -> Extraction:
    result = Extraction(name=name, sha256=sha256(raw).hexdigest())
    total = 0

    def add(locator, text):
        nonlocal total
        text = str(text).replace("\x00", "")
        if not text.strip():
            return
        total += len(text)
        if total > settings.max_chars:
            raise IngestionError("extracted_text_limit")
        # Split long passages without discarding text; offsets are extraction offsets,
        # not invented page/paragraph coordinates.
        for offset in range(0, len(text), 6000):
            loc = (
                locator
                if len(text) <= 6000
                else f"{locator}; characters {offset + 1}-{min(offset + 6000, len(text))}"
            )
            cid = sha256(f"{result.sha256}:{name}:{loc}".encode()).hexdigest()[:24]
            result.chunks.append(
                Chunk(cid, result.sha256, name, loc, text[offset : offset + 6000])
            )

    try:
        if len(raw) > settings.max_file_bytes:
            raise IngestionError("file_size_limit")
        ext = Path(name).suffix.lower()
        if ext not in SUPPORTED:
            raise IngestionError("unsupported_format")
        if ext in {".docx", ".xlsx"}:
            check_package(raw)
        if ext == ".txt":
            lines = raw.decode("utf-8-sig").splitlines()
            for i in range(0, len(lines), 30):
                add(
                    f"lines {i + 1}-{min(i + 30, len(lines))}",
                    "\n".join(lines[i : i + 30]),
                )
        elif ext == ".pdf":
            pdf = PdfReader(BytesIO(raw))
            if pdf.is_encrypted:
                raise IngestionError("encrypted_pdf")
            if len(pdf.pages) > settings.max_pages:
                raise IngestionError("page_limit")
            empty = []
            for n, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ""
                if not text.strip():
                    empty.append(n)
                add(f"page {n}", text)
            if empty:
                result.warnings.append(
                    "OCR needed or blank pages; unextracted pages: "
                    + ", ".join(map(str, empty))
                )
                result.status = "ocr_needed"
        elif ext == ".docx":
            doc = Document(BytesIO(raw))
            for n, p in enumerate(doc.paragraphs, 1):
                add(f"paragraph {n}", p.text)

            def table_cells(t, prefix):
                for r, row in enumerate(t.rows, 1):
                    for c, cell in enumerate(row.cells, 1):
                        add(f"{prefix} row {r} cell {c}", cell.text)
                        for i, nested in enumerate(cell.tables, 1):
                            table_cells(
                                nested, f"{prefix} row {r} cell {c} nested table {i}"
                            )

            for t, table in enumerate(doc.tables, 1):
                table_cells(table, f"table {t}")
            for s, section in enumerate(doc.sections, 1):
                for label, part in [
                    ("header", section.header),
                    ("footer", section.footer),
                ]:
                    for n, p in enumerate(part.paragraphs, 1):
                        add(f"section {s} {label} paragraph {n}", p.text)
            with zipfile.ZipFile(BytesIO(raw)) as z:
                xml = z.read("word/document.xml")
                if re.search(rb"<w:(ins|del|txbxContent)\b", xml):
                    result.warnings.append(
                        "Tracked changes or text boxes require manual review; extraction may omit this content."
                    )
                    result.status = "needs_review"
                for part in [
                    "word/footnotes.xml",
                    "word/endnotes.xml",
                    "word/comments.xml",
                ]:
                    if part in z.namelist():
                        from defusedxml import ElementTree

                        root = ElementTree.fromstring(z.read(part))
                        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
                        for n, node in enumerate(root, 1):
                            add(
                                f"{Path(part).stem} item {n}",
                                " ".join(t.text or "" for t in node.iter(ns + "t")),
                            )
        elif ext == ".xlsx":
            workbook = load_workbook(
                BytesIO(raw), read_only=True, data_only=False, keep_links=False
            )
            values = load_workbook(
                BytesIO(raw), read_only=True, data_only=True, keep_links=False
            )
            cells_seen = 0
            try:
                for sheet in workbook:
                    if (sheet.max_row or 0) * (
                        sheet.max_column or 0
                    ) > settings.max_cells:
                        raise IngestionError("cell_limit")
                    for row, cached_row in zip(
                        sheet.iter_rows(), values[sheet.title].iter_rows(), strict=True
                    ):
                        for cell, cached in zip(row, cached_row, strict=True):
                            cells_seen += 1
                            if cells_seen > settings.max_cells:
                                raise IngestionError("cell_limit")
                            if cell.value is None:
                                continue
                            val = str(cell.value)
                            if cell.data_type == "f":
                                val = f"formula: {cell.value}; cached value: {cached.value if cached.value is not None else 'UNAVAILABLE'}"
                                if cached.value is None:
                                    result.warnings.append(
                                        f"{sheet.title}!{cell.coordinate}: formula cached value unavailable; not evaluated."
                                    )
                            add(f"sheet {sheet.title!r} cell {cell.coordinate}", val)
            finally:
                workbook.close()
                values.close()
        elif ext == ".eml":
            msg = BytesParser(policy=policy.default).parsebytes(raw)
            for key in [
                "From",
                "To",
                "Cc",
                "Bcc",
                "Subject",
                "Date",
                "Message-ID",
                "In-Reply-To",
            ]:
                for n, value in enumerate(msg.get_all(key, []), 1):
                    add(f"header {key} [{n}]", str(value))
            count = 0
            for part in msg.walk():
                count += 1
                if count > 200:
                    raise IngestionError("email_part_limit")
                if (
                    part.get_filename()
                    or part.get_content_disposition() == "attachment"
                ):
                    result.warnings.append(
                        f"Attachment {part.get_filename() or '(unnamed)'}: not ingested; upload separately. Original retained in EML."
                    )
                    continue
                if part.is_multipart():
                    continue
                if part.get_content_type() in {"text/plain", "text/html"}:
                    body = part.get_content()
                    if part.get_content_type() == "text/html":
                        soup = BeautifulSoup(body, "html.parser")
                        for tag in soup(["script", "style"]):
                            tag.decompose()
                        body = soup.get_text("\n")
                    add(f"body MIME segment {count}", body)
                else:
                    result.warnings.append(
                        f"MIME segment {count}: unsupported {part.get_content_type()}; original retained."
                    )
        if not result.chunks and result.status == "ready":
            result.status = "empty"
            result.error_code = "no_extractable_text"
    except IngestionError as e:
        result.status, result.error_code = "failed", str(e)
        result.chunks = []
    except Exception:
        result.status, result.error_code = "failed", "invalid_or_unreadable_file"
        result.chunks = []
    return result


def search(chunks: list[dict], query: str, limit=30):
    terms = set(re.findall(r"[\w'-]+", query.lower()))
    ranked = sorted(
        chunks,
        key=lambda c: sum(c["text"].lower().count(t) for t in terms),
        reverse=True,
    )
    return ranked[:limit]
